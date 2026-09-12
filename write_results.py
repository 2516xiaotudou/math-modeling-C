# -*- coding: utf-8 -*-
"""生成 result1/2/3/4-2/4-3.xlsx，匹配附件5模板格式。"""
import os
from datetime import datetime, timedelta
import numpy as np
import pandas as pd
from openpyxl import Workbook

BASE = r'D:\桌面\C题\附件'
OUTDIR = r'D:\桌面\C题\结果'
os.makedirs(OUTDIR, exist_ok=True)
N = 144
OFF = 31          # 结果从 2025-2-1(全年第31天,0基) 开始
ND = 334          # 2.1 ~ 12.31 共334天

# ---- 载入中间结果 ----
d = np.load(os.path.join(OUTDIR, '_data.npz'))
b1 = d['b1']; c1 = d['c1']; d1 = d['d1']; SOC1 = d['SOC1']
b2_all = d['b2_all'][OFF:]; c2_all = d['c2_all'][OFF:]; d2_all = d['d2_all'][OFF:]; SOC2_traj = d['SOC2_traj'][OFF:]
b_plan3 = d['b_plan3'][OFF:]; b_adj3 = d['b_adj3'][OFF:]
c3_all = d['c3_all'][OFF:]; d3_all = d['d3_all'][OFF:]; e3_all = d['e3_all'][OFF:]; SOC3_traj = d['SOC3_traj'][OFF:]
b42_all = d['b42_all'][OFF:]; c42_all = d['c42_all'][OFF:]; d42_all = d['d42_all'][OFF:]; SOC42 = d['SOC42'][OFF:]
b_plan43 = d['b_plan43'][OFF:]; b_adj43 = d['b_adj43'][OFF:]
c43_all = d['c43_all'][OFF:]; d43_all = d['d43_all'][OFF:]; e43_all = d['e43_all'][OFF:]; SOC43 = d['SOC43'][OFF:]

df1 = pd.read_excel(os.path.join(BASE, '附件1.xlsx'))
p1 = df1['电价'].values.astype(float)
df4 = pd.read_excel(os.path.join(BASE, '附件4.xlsx'))
price_all = df4.iloc[:, 1:].values.astype(float)[OFF:]

dates = [datetime(2025, 1, 1) + timedelta(days=i) for i in range(OFF, OFF + ND)]

# ---- 时间标签 ----
def fm(t):                                   # 分钟 -> "HH:MM"(>=24:00 加 +1)
    carry = t // 1440
    m = t % 1440
    s = '%02d:%02d' % (m // 60, m % 60)
    return s + ('+1' if carry else '')

LL = ['%s-%s' % (fm((j + 1) * 10), fm((j + 2) * 10)) for j in range(N)]   # result1 长表标签
LW = list(LL)
LW[143] = '0:00-0:10+1'                      # 宽表末标签(模板差异)

def rot(x):                                   # 数据区间 -> 模板列(旋转:列j=区间(j+1)%144)
    return np.roll(x, -1)

PERIODS = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00', '16:00-20:00', '20:00-24:00']

def write_wide(ws, b, tot_cost):
    ws.append(['日期\\时间'] + LW + ['全天购电量', '全天购电费'])
    for di in range(ND):
        br = rot(b[di])
        ws.append([dates[di]] + [round(float(x), 4) for x in br] +
                  [round(float(br.sum()), 4), round(float(tot_cost[di]), 4)])

def write_cd(ws, c_all, d_all, soc_traj, has_date=True):
    ws.append((['日期'] if has_date else []) + ['时间段', '充电量', '放电量', '时刻', '储电量'])
    for di in range(ND):
        c = c_all[di]; dd = d_all[di]; soc = soc_traj[di]
        for k in range(6):
            chg = round(float(c[24 * k:24 * k + 24].sum()), 4)
            dis = round(float(dd[24 * k:24 * k + 24].sum()), 4)
            tms = '0:00' if k == 0 else ('24:00' if k == 1 else None)
            stg = round(float(soc[0]), 4) if k == 0 else (round(float(soc[144]), 4) if k == 1 else None)
            base = [dates[di] if k == 0 else None] if has_date else []
            ws.append(base + [PERIODS[k], chg, dis, tms, stg])

def merge_events(e):
    ev = []
    t = 0
    while t < N:
        if e[t] > 1e-6:
            s = t
            while t < N and e[t] > 1e-6:
                t += 1
            ev.append((s, t - 1, float(e[s:t].sum())))
        else:
            t += 1
    return ev

def write_emergency(ws, e_all):
    ws.append(['日期', '购电时间段', '购电量'])
    for di in range(ND):
        ev = merge_events(e_all[di])
        if not ev:
            continue
        ws.append([dates[di], None, None])
        for s, t, en in ev:
            ws.append([None, '%s-%s' % (fm(s * 10), fm((t + 1) * 10)), round(en, 4)])

# ---- 每日购电费 ----
cost2_d = np.array([float(p1 @ b2_all[di]) for di in range(ND)])

plan3_d = np.zeros(ND); adj3_d = np.zeros(ND); emerg3_d = np.zeros(ND)
for di in range(ND):
    bp = b_plan3[di]; ba = b_adj3[di]; e = e3_all[di]
    mn = np.minimum(bp, ba); over = np.clip(bp - ba, 0, None); under = np.clip(ba - bp, 0, None)
    plan3_d[di] = float(p1 @ mn)
    adj3_d[di] = float(p1 @ (0.5 * over + 1.5 * under))
    emerg3_d[di] = float(5.0 * (p1 @ e))
tot3_d = plan3_d + adj3_d + emerg3_d

cost42_d = np.array([float(price_all[di] @ b42_all[di]) for di in range(ND)])

plan43_d = np.zeros(ND); adj43_d = np.zeros(ND); emerg43_d = np.zeros(ND)
for di in range(ND):
    p = price_all[di]
    bp = b_plan43[di]; ba = b_adj43[di]; e = e43_all[di]
    mn = np.minimum(bp, ba); over = np.clip(bp - ba, 0, None); under = np.clip(ba - bp, 0, None)
    plan43_d[di] = float(p @ mn)
    adj43_d[di] = float(p @ (0.5 * over + 1.5 * under))
    emerg43_d[di] = float(5.0 * (p @ e))
tot43_d = plan43_d + adj43_d + emerg43_d

zero_e = np.zeros((ND, N))

# ---- result1 (2 sheets) ----
wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
ws.append(['时间段', '购电量'])
br1 = rot(b1)
for j in range(N):
    ws.append([LL[j], round(float(br1[j]), 4)])
ws2 = wb.create_sheet('充放电量')
ws2.append(['时间段', '充电量', '放电量', '时刻', '储电量'])
for k in range(6):
    chg = round(float(c1[24 * k:24 * k + 24].sum()), 4)
    dis = round(float(d1[24 * k:24 * k + 24].sum()), 4)
    tms = '0:00' if k == 0 else ('24:00' if k == 1 else None)
    stg = round(float(SOC1[0]), 4) if k == 0 else (round(float(SOC1[144]), 4) if k == 1 else None)
    ws2.append([PERIODS[k], chg, dis, tms, stg])
wb.save(os.path.join(OUTDIR, 'result1.xlsx'))

# ---- result2 (3 sheets) ----
wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
write_wide(ws, b2_all, cost2_d)
ws2 = wb.create_sheet('充放电量'); write_cd(ws2, c2_all, d2_all, SOC2_traj, True)
ws3 = wb.create_sheet('紧急购电量'); write_emergency(ws3, zero_e)
wb.save(os.path.join(OUTDIR, 'result2.xlsx'))

# ---- result3 (4 sheets) ----
wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
write_wide(ws, b_plan3, tot3_d)
ws2 = wb.create_sheet('调整购电量'); write_wide(ws2, b_adj3, tot3_d)
ws3 = wb.create_sheet('充放电量'); write_cd(ws3, c3_all, d3_all, SOC3_traj, True)
ws4 = wb.create_sheet('紧急购电量'); write_emergency(ws4, e3_all)
wb.save(os.path.join(OUTDIR, 'result3.xlsx'))

# ---- result4-2 (3 sheets) ----
wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
write_wide(ws, b42_all, cost42_d)
ws2 = wb.create_sheet('充放电量'); write_cd(ws2, c42_all, d42_all, SOC42, True)
ws3 = wb.create_sheet('紧急购电量'); write_emergency(ws3, zero_e)
wb.save(os.path.join(OUTDIR, 'result4-2.xlsx'))

# ---- result4-3 (4 sheets) ----
wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
write_wide(ws, b_plan43, tot43_d)
ws2 = wb.create_sheet('调整购电量'); write_wide(ws2, b_adj43, tot43_d)
ws3 = wb.create_sheet('充放电量'); write_cd(ws3, c43_all, d43_all, SOC43, True)
ws4 = wb.create_sheet('紧急购电量'); write_emergency(ws4, e43_all)
wb.save(os.path.join(OUTDIR, 'result4-3.xlsx'))

# ---- 校验汇总 ----
ver = []
ver.append('result1 total cost %.2f, purchase %.2f' % (float(p1 @ b1), float(b1.sum())))
ver.append('result2 total cost %.2f' % cost2_d.sum())
ver.append('result3 plan %.2f adj %.2f emerg %.2f total %.2f' %
           (plan3_d.sum(), adj3_d.sum(), emerg3_d.sum(), tot3_d.sum()))
ver.append('result3 emergency energy %.2f kWh' % e3_all.sum())
ver.append('result4-2 total cost %.2f' % cost42_d.sum())
ver.append('result4-3 plan %.2f adj %.2f emerg %.2f total %.2f' %
           (plan43_d.sum(), adj43_d.sum(), emerg43_d.sum(), tot43_d.sum()))
ver.append('result4-3 emergency energy %.2f kWh' % e43_all.sum())
with open(os.path.join(OUTDIR, '_verify.txt'), 'w', encoding='utf-8') as f:
    f.write('\n'.join(ver))
print('WRITE DONE')
