import os
from datetime import datetime, timedelta
import numpy as np
import pandas as pd
from openpyxl import Workbook

bd = r'D:\桌面\C题\附件'
od = r'D:\桌面\C题\结果'
os.makedirs(od, exist_ok=True)
n = 144
of = 31
nd = 334

z = np.load(os.path.join(od, '_data.npz'))
b1 = z['b1']; c1 = z['c1']; d1 = z['d1']; s1 = z['s1']
b2 = z['b2'][of:]; c2 = z['c2'][of:]; d2 = z['d2'][of:]; s2 = z['s2'][of:]
bp3 = z['bp3'][of:]; ba3 = z['ba3'][of:]
c3 = z['c3'][of:]; d3 = z['d3'][of:]; e3 = z['e3'][of:]; s3 = z['s3'][of:]
b42 = z['b42'][of:]; c42 = z['c42'][of:]; d42 = z['d42'][of:]; s42 = z['s42'][of:]
bp43 = z['bp43'][of:]; ba43 = z['ba43'][of:]
c43 = z['c43'][of:]; d43 = z['d43'][of:]; e43 = z['e43'][of:]; s43 = z['s43'][of:]

a1 = pd.read_excel(os.path.join(bd, '附件1.xlsx'))
pr = a1['电价'].values.astype(float)
a5 = pd.read_excel(os.path.join(bd, '附件4.xlsx'))
pc = a5.iloc[:, 1:].values.astype(float)[of:]

ds = [datetime(2025, 1, 1) + timedelta(days=i) for i in range(of, of + nd)]

def f(t):
    cr = t // 1440
    m = t % 1440
    s = '%02d:%02d' % (m // 60, m % 60)
    return s + ('+1' if cr else '')

l1 = ['%s-%s' % (f((j + 1) * 10), f((j + 2) * 10)) for j in range(n)]
l2 = list(l1)
l2[143] = '0:00-0:10+1'

def rt(x):
    return np.roll(x, -1)

ps = ['0:00-4:00', '4:00-8:00', '8:00-12:00', '12:00-16:00', '16:00-20:00', '20:00-24:00']

def ww(ws, b, tc):
    ws.append(['日期\\时间'] + l2 + ['全天购电量', '全天购电费'])
    for di in range(nd):
        br = rt(b[di])
        ws.append([ds[di]] + [round(float(x), 4) for x in br] +
                  [round(float(br.sum()), 4), round(float(tc[di]), 4)])

def wc(ws, ca, da, st, hd=True):
    ws.append((['日期'] if hd else []) + ['时间段', '充电量', '放电量', '时刻', '储电量'])
    for di in range(nd):
        c = ca[di]; dd = da[di]; ss = st[di]
        for k in range(6):
            ch = round(float(c[24 * k:24 * k + 24].sum()), 4)
            dg = round(float(dd[24 * k:24 * k + 24].sum()), 4)
            tm = '0:00' if k == 0 else ('24:00' if k == 1 else None)
            sg = round(float(ss[0]), 4) if k == 0 else (round(float(ss[144]), 4) if k == 1 else None)
            base = [ds[di] if k == 0 else None] if hd else []
            ws.append(base + [ps[k], ch, dg, tm, sg])

def me(e):
    ev = []
    t = 0
    while t < n:
        if e[t] > 1e-6:
            s = t
            while t < n and e[t] > 1e-6:
                t += 1
            ev.append((s, t - 1, float(e[s:t].sum())))
        else:
            t += 1
    return ev

def we(ws, ea):
    ws.append(['日期', '购电时间段', '购电量'])
    for di in range(nd):
        ev = me(ea[di])
        if not ev:
            continue
        ws.append([ds[di], None, None])
        for s, t, en in ev:
            ws.append([None, '%s-%s' % (f(s * 10), f((t + 1) * 10)), round(en, 4)])

ct2 = np.array([float(pr @ b2[di]) for di in range(nd)])

cp3 = np.zeros(nd); ca3 = np.zeros(nd); ce3 = np.zeros(nd)
for di in range(nd):
    bpa = bp3[di]; baa = ba3[di]; e = e3[di]
    mn = np.minimum(bpa, baa); ov = np.clip(bpa - baa, 0, None); un = np.clip(baa - bpa, 0, None)
    cp3[di] = float(pr @ mn)
    ca3[di] = float(pr @ (0.5 * ov + 1.5 * un))
    ce3[di] = float(5.0 * (pr @ e))
tt3 = cp3 + ca3 + ce3

ct42 = np.array([float(pc[di] @ b42[di]) for di in range(nd)])

cp43 = np.zeros(nd); ca43 = np.zeros(nd); ce43 = np.zeros(nd)
for di in range(nd):
    p = pc[di]
    bpa = bp43[di]; baa = ba43[di]; e = e43[di]
    mn = np.minimum(bpa, baa); ov = np.clip(bpa - baa, 0, None); un = np.clip(baa - bpa, 0, None)
    cp43[di] = float(p @ mn)
    ca43[di] = float(p @ (0.5 * ov + 1.5 * un))
    ce43[di] = float(5.0 * (p @ e))
tt43 = cp43 + ca43 + ce43

ze = np.zeros((nd, n))

wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
ws.append(['时间段', '购电量'])
br1 = rt(b1)
for j in range(n):
    ws.append([l1[j], round(float(br1[j]), 4)])
ws2 = wb.create_sheet('充放电量')
ws2.append(['时间段', '充电量', '放电量', '时刻', '储电量'])
for k in range(6):
    ch = round(float(c1[24 * k:24 * k + 24].sum()), 4)
    dg = round(float(d1[24 * k:24 * k + 24].sum()), 4)
    tm = '0:00' if k == 0 else ('24:00' if k == 1 else None)
    sg = round(float(s1[0]), 4) if k == 0 else (round(float(s1[144]), 4) if k == 1 else None)
    ws2.append([ps[k], ch, dg, tm, sg])
wb.save(os.path.join(od, 'result1.xlsx'))

wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
ww(ws, b2, ct2)
ws2 = wb.create_sheet('充放电量'); wc(ws2, c2, d2, s2, True)
ws3 = wb.create_sheet('紧急购电量'); we(ws3, ze)
wb.save(os.path.join(od, 'result2.xlsx'))

wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
ww(ws, bp3, cp3)
ws2 = wb.create_sheet('调整购电量'); ww(ws2, ba3, ca3)
ws3 = wb.create_sheet('充放电量'); wc(ws3, c3, d3, s3, True)
ws4 = wb.create_sheet('紧急购电量'); we(ws4, e3)
wb.save(os.path.join(od, 'result3.xlsx'))

wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
ww(ws, b42, ct42)
ws2 = wb.create_sheet('充放电量'); wc(ws2, c42, d42, s42, True)
ws3 = wb.create_sheet('紧急购电量'); we(ws3, ze)
wb.save(os.path.join(od, 'result4-2.xlsx'))

wb = Workbook(); ws = wb.active; ws.title = '计划购电量'
ww(ws, bp43, cp43)
ws2 = wb.create_sheet('调整购电量'); ww(ws2, ba43, ca43)
ws3 = wb.create_sheet('充放电量'); wc(ws3, c43, d43, s43, True)
ws4 = wb.create_sheet('紧急购电量'); we(ws4, e43)
wb.save(os.path.join(od, 'result4-3.xlsx'))

vv = []
vv.append('r1 cost %.2f buy %.2f' % (float(pr @ b1), float(b1.sum())))
vv.append('r2 cost %.2f' % ct2.sum())
vv.append('r3 plan %.2f adj %.2f em %.2f total %.2f' % (cp3.sum(), ca3.sum(), ce3.sum(), tt3.sum()))
vv.append('r3 em energy %.2f' % e3.sum())
vv.append('r4-2 cost %.2f' % ct42.sum())
vv.append('r4-3 plan %.2f adj %.2f em %.2f total %.2f' % (cp43.sum(), ca43.sum(), ce43.sum(), tt43.sum()))
vv.append('r4-3 em energy %.2f' % e43.sum())
with open(os.path.join(od, '_verify.txt'), 'w', encoding='utf-8') as fh:
    fh.write('\n'.join(vv))
print('WRITE DONE')
